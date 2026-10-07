import io
import os
import sys
from pathlib import Path

import pytest

import tempfile

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
os.environ.setdefault("APP_DATA_DIR", tempfile.mkdtemp(prefix="boq-test-"))

from app import boq_engine as eng  # noqa: E402


def el(**kw):
    return eng.Element.from_dict(kw)


def qty(lines, code):
    return sum(ln.qty for ln in lines if ln.code == code)


def test_column_quantities():
    lines = eng.takeoff_element(el(type="COT", code="C1", n=6,
                                   params={"b": 0.3, "h": 0.4, "H": 3.6, "ham_luong": 150}))
    assert qty(lines, "BT_COT") == pytest.approx(6 * 0.3 * 0.4 * 3.6)
    assert qty(lines, "VK_COT") == pytest.approx(6 * 2 * (0.3 + 0.4) * 3.6)
    assert qty(lines, "CT_COT") == pytest.approx(6 * 0.3 * 0.4 * 3.6 * 150 / 1000)
    assert lines[0].name == "Bê tông cột B22.5"


def test_beam_deducts_slab():
    lines = eng.takeoff_element(el(type="DAM", n=2, grade="B25",
                                   params={"b": 0.22, "h": 0.5, "L": 6, "t_san": 0.12}))
    assert qty(lines, "BT_DAM") == pytest.approx(2 * 0.22 * 0.38 * 6)
    assert qty(lines, "VK_DAM") == pytest.approx(2 * (0.22 + 2 * 0.38) * 6)
    assert lines[0].item == "BT_DAM|B25"


def test_footing_wall_floor_and_custom():
    m = eng.takeoff_element(el(type="MONG", n=4, params={"L": 1.5, "B": 1.2, "H": 0.5, "lot": 0.1,
                                                         "sau_dao": 1.5, "mo_rong": 0.3}))
    assert qty(m, "DAO") == pytest.approx(4 * 2.1 * 1.8 * 1.5)
    assert qty(m, "BTL") == pytest.approx(4 * 1.7 * 1.4 * 0.1)
    w = eng.takeoff_element(el(type="TUONG", params={"L": 6, "H": 3, "t": 0.2, "lo_cua": 2, "mat_trat": 2}))
    assert qty(w, "XAY") == pytest.approx((18 - 2) * 0.2)
    assert qty(w, "TRAT") == pytest.approx(32)
    n = eng.takeoff_element(el(type="NEN", params={"L": 5, "B": 4}))
    assert qty(n, "LAT") == pytest.approx(20)
    k = eng.takeoff_element(el(type="KHAC", n=3, item_name="Cửa D1", unit="bộ", params={"khoi_luong": 1}))
    assert k[0].qty == 3 and k[0].unit == "bộ"


def test_every_expression_matches_quantity():
    elements = [el(type=t, n=3) for t in eng.ELEMENT_TYPES if t != "KHAC"]
    for ln in eng.takeoff(elements):
        assert eng.eval_expr(ln.expr) == pytest.approx(ln.qty)
        assert ln.expr.startswith("3*")


def test_invalid_input():
    with pytest.raises(eng.BoqError):
        el(type="XYZ")
    with pytest.raises(eng.BoqError):
        el(type="COT", params={"b": "abc"})
    with pytest.raises(eng.BoqError):
        el(type="COT", n=0)
    with pytest.raises(eng.BoqError):
        eng.takeoff_element(el(type="DAM", params={"h": 0.1, "t_san": 0.12}))
    with pytest.raises(eng.BoqError):
        eng.eval_expr("__import__('os')")
    assert el(type="COT", params={"b": "0,25"}).params["b"] == 0.25


def test_summary_groups_and_prices():
    lines = eng.takeoff([el(type="COT", n=2), el(type="COT", n=3, code="C2")])
    summary = eng.summarize(lines, {"BT_COT|B22.5": 1_000_000})
    bt = next(s for s in summary if s["code"] == "BT_COT")
    assert bt["qty"] == pytest.approx(5 * 0.3 * 0.3 * 3.6)
    assert bt["amount"] == round(bt["qty"] * 1_000_000)
    assert [s["code"] for s in summary] == ["BT_COT", "VK_COT", "CT_COT"]


def test_statistics_and_3d():
    elements = [el(type="COT", n=4, dx=6, floor="L01"), el(type="SAN", floor="L01"),
                el(type="KHAC", params={"khoi_luong": 1})]
    s = eng.statistics(elements)
    assert s["totals"]["elements"] == 6
    assert s["by_floor"][0]["floor"] == "L01"
    m = eng.model3d(elements)
    assert len(m["boxes"]) == 5  # 4 columns + 1 slab, KHAC has no geometry
    assert [b["c"][0] for b in m["boxes"][:4]] == [0, 6, 12, 18]


def test_shop_svg():
    svg = eng.shop_svg(el(type="DAM", code="D<1>", params={"b": 0.22, "h": 0.5, "L": 6}))
    assert svg.startswith("<svg") and "D&lt;1&gt;" in svg and "220" in svg and "6000" in svg


def _dxf_bytes():
    import ezdxf
    doc = ezdxf.new()
    msp = doc.modelspace()
    for i in range(3):  # 3 columns 300x400 every 6000 mm
        x = i * 6000
        msp.add_lwpolyline([(x, 0), (x + 300, 0), (x + 300, 400), (x, 400)], close=True,
                           dxfattribs={"layer": "KC-COT"})
    msp.add_lwpolyline([(0, 0), (12000, 0), (12000, 220), (0, 220)], close=True,
                       dxfattribs={"layer": "KC-DẦM"})
    msp.add_lwpolyline([(0, 0), (10, 0), (10, 10)], close=True, dxfattribs={"layer": "TEXT"})
    buf = io.StringIO()
    doc.write(buf)
    return buf.getvalue().encode()


def test_dxf_import():
    res = eng.import_dxf(_dxf_bytes(), unit_scale=0.001, floor="L01", defaults={"H_cot": 3.3})
    cols = [e for e in res["elements"] if e["type"] == "COT"]
    beams = [e for e in res["elements"] if e["type"] == "DAM"]
    assert len(cols) == 1 and cols[0]["n"] == 3 and cols[0]["dx"] == 6.0
    assert cols[0]["params"] == {"b": 0.3, "h": 0.4, "H": 3.3, "ham_luong": 150}
    assert beams[0]["params"]["L"] == 12.0 and beams[0]["params"]["b"] == 0.22
    assert beams[0]["params"]["t_san"] == 0.12
    assert beams[0]["z"] == pytest.approx(3.3 - 0.5)  # beam top flush with column top
    assert res["skipped_layers"] == {"TEXT": 1}
    for cand in res["elements"]:
        eng.Element.from_dict(cand)


def test_template_roundtrip():
    res = eng.import_template(eng.template_workbook())
    assert not res["errors"]
    assert [e["type"] for e in res["elements"]] == list(eng.ELEMENT_TYPES)
    beam = next(e for e in res["elements"] if e["type"] == "DAM")
    assert beam["params"]["t_san"] == 0.12 and beam["n"] == 5


def test_export_workbook_formulas():
    from openpyxl import load_workbook
    elements = [el(type="COT", code="C1", n=6, floor="L01"), el(type="DAM", code="D1", n=2)]
    data = eng.export_workbook({"name": "Test", "code": "T"}, elements, {"BT_COT|B22.5": 1500000})
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["TongHop_BOQ", "ChiTiet", "ThongKe_CauKien"]
    detail = wb["ChiTiet"]
    first = detail.cell(row=5, column=10).value
    assert first.startswith("=6*") and eng.eval_expr(first[1:]) == pytest.approx(6 * 0.3 * 0.3 * 3.6)
    summary = wb["TongHop_BOQ"]
    values = [summary.cell(row=r, column=5).value for r in range(5, summary.max_row + 1)]
    assert any(str(v).startswith("=SUMIF(ChiTiet!") for v in values)


# ---------------------------------------------------------------- API (FastAPI)

@pytest.fixture()
def client():
    from fastapi.testclient import TestClient
    from app.main import app
    return TestClient(app)


def test_api_end_to_end(client):
    p = client.post("/api/v1/boq/projects", json={"name": "Nhà xưởng", "code": "NX-01"}).json()
    pid = p["id"]
    r = client.post(f"/api/v1/boq/projects/{pid}/import/dxf", data={"unit": "mm", "floor": "L01"},
                    files={"file": ("plan.dxf", _dxf_bytes(), "application/dxf")})
    assert r.status_code == 200, r.text
    cands = r.json()["elements"]
    r = client.post(f"/api/v1/boq/projects/{pid}/elements", json={"elements": cands})
    assert r.json()["added"] == 2
    bad = client.post(f"/api/v1/boq/projects/{pid}/elements", json={"elements": [{"type": "COT", "n": -1}]})
    assert bad.status_code == 400
    proj = client.get(f"/api/v1/boq/projects/{pid}").json()
    eid = proj["elements"][0]["id"]
    assert client.get(f"/api/v1/boq/elements/{eid}/shop.svg").text.startswith("<svg")
    assert len(client.get(f"/api/v1/boq/projects/{pid}/model3d").json()["boxes"]) == 4
    assert client.get(f"/api/v1/boq/projects/{pid}/stats").json()["totals"]["elements"] == 4
    client.put(f"/api/v1/boq/projects/{pid}/prices", json={"prices": {"BT_COT|B22.5": 1200000}})
    t = client.get(f"/api/v1/boq/projects/{pid}/takeoff").json()
    assert next(s for s in t["summary"] if s["code"] == "BT_COT")["price"] == 1200000
    x = client.get(f"/api/v1/boq/projects/{pid}/export.xlsx")
    assert x.status_code == 200 and x.content[:2] == b"PK"
    assert client.get("/boq").status_code == 200
    assert client.delete(f"/api/v1/boq/elements/{eid}").status_code == 200


# ---------------------------------------------------------------- WBS + dự toán

from app import estimate_engine as est  # noqa: E402


def test_wbs_codes():
    els = [el(type="COT", code="C1", hang_muc="Nhà A", id=1), el(type="COT", code="C2", hang_muc="Nhà A", id=2),
           el(type="MONG", code="M1", hang_muc="Nhà A", id=3), el(type="COT", code="C9", hang_muc="Nhà B", id=4)]
    lines = eng.takeoff(els)
    bt = [ln for ln in lines if ln.code == "BT_COT"]
    # Nhà A = 01, Kết cấu is the 2nd phần việc after Phần ngầm, BT cột first công tác.
    assert [ln.wbs for ln in bt] == ["01.02.01.001", "01.02.01.002", "02.01.01.001"]
    assert lines[0].wbs.startswith("01.01.")  # Phần ngầm sorts first
    summary = eng.summarize(lines)
    assert [s["wbs"] for s in summary if s["code"] == "BT_COT"] == ["01.02.01", "02.01.01"]
    nodes = {n["code"]: n for n in eng.wbs_tree(lines)}
    assert nodes["01"]["level"] == 100 and nodes["01"]["name"] == "Nhà A"
    assert nodes["01.02"]["level"] == 300 and nodes["01.02"]["name"] == "Kết cấu"
    assert nodes["01.02.01"]["qty"] == pytest.approx(2 * 0.3 * 0.3 * 3.6)


def test_unit_factor():
    assert est.unit_factor("m2", "100m2") == pytest.approx(0.01)
    assert est.unit_factor("m3", "m³") == 1
    assert est.unit_factor("tấn", "T") == 1
    assert est.unit_factor("m2", "m3") is None


def _estimate(els, lib=None, direct=None, rates=None):
    return est.estimate(eng.summarize(eng.takeoff(els)), lib or est.sample_library(), direct, rates)


def test_estimate_unit_price_and_costs():
    lib = est.Library()
    lib.resources["XM"] = est.Resource("XM", "Xi măng", "kg", "VL", 2000)
    lib.resources["NC"] = est.Resource("NC", "Nhân công", "công", "NC", 300000)
    lib.resources["MAY"] = est.Resource("MAY", "Máy trộn", "ca", "M", 400000)
    lib.norms["DM1"] = est.Norm("DM1", "BT cột", "m3", [("XM", 350), ("NC", 3), ("MAY", 0.1)])
    lib.mapping["BT_COT"] = "DM1"
    e = _estimate([el(type="COT", n=10, params={"ham_luong": 0})], lib,
                  rates={"C": 6, "LT": 1, "TT": 2, "TL": 5, "VAT": 10, "DP": 5})
    item = e["items"][0]
    qty = 10 * 0.3 * 0.3 * 3.6
    assert item["dg"] == pytest.approx({"VL": 700000, "NC": 900000, "M": 40000, "K": 0})
    assert item["total"] == pytest.approx(qty * 1640000)
    costs = {c["key"]: c["value"] for c in e["costs"]}
    T = qty * 1640000
    assert costs["T"] == pytest.approx(T)
    assert costs["GT"] == pytest.approx(T * 0.09)
    assert costs["TL"] == pytest.approx(T * 1.09 * 0.05)
    G = T * 1.09 * 1.05
    assert costs["G"] == pytest.approx(G)
    assert costs["TONG"] == pytest.approx(G * 1.1 * 1.05)
    res = {r["code"]: r for r in e["resources"]}
    assert res["XM"]["haophi"] == pytest.approx(qty * 350)
    # VK_COT has no norm and no direct price → warning
    assert any("Ván khuôn cột" in w for w in e["warnings"])


def test_estimate_unit_conversion_and_direct_price():
    lib = est.sample_library()
    norm = lib.norms["MAU.VK.SAN"]
    per_m2 = _estimate([el(type="SAN")], lib)["items"]
    norm.unit = "100m2"
    norm.items = [(r, q * 100) for r, q in norm.items]
    per_100 = _estimate([el(type="SAN")], lib)["items"]
    vk = lambda items: next(i for i in items if i["code"] == "VK_SAN")  # noqa: E731
    assert vk(per_100)["factor"] == pytest.approx(0.01)
    assert vk(per_100)["unit_price"] == pytest.approx(vk(per_m2)["unit_price"])
    norm.unit = "m3"  # wrong unit → not priced, warned
    e = _estimate([el(type="SAN")], lib)
    assert vk(e["items"])["norm"] is None and any("khác đơn vị" in w for w in e["warnings"])
    k = _estimate([el(type="KHAC", n=2, item_name="Cửa", unit="bộ", params={"khoi_luong": 1})],
                  lib, {"KHAC|Cửa|bộ": 5000000})
    assert k["direct"]["K"] == 10000000


def test_library_roundtrip():
    lib = est.sample_library()
    back, errors = est.import_library(est.library_workbook(lib))
    assert not errors
    assert back.to_dict() == lib.to_dict()


def test_estimate_excel_recalculates():
    from openpyxl import load_workbook
    from pycel import ExcelCompiler
    els = [el(type=t, n=2, hang_muc="Nhà A") for t in eng.ELEMENT_TYPES if t != "KHAC"]
    e = _estimate(els, rates={"VAT": 8, "DP": 5})
    assert not e["warnings"]
    path = Path(tempfile.mkdtemp()) / "dt.xlsx"
    path.write_bytes(est.export_estimate({"name": "T"}, e))
    ws = load_workbook(path)["TongHop_ChiPhi"]
    row = {ws.cell(r, 3).value: r for r in range(5, 25) if ws.cell(r, 3).value}
    xl = ExcelCompiler(filename=str(path))
    for c in e["costs"]:
        assert xl.evaluate(f"TongHop_ChiPhi!F{row[c['key']]}") == pytest.approx(c["value"], abs=0.5)
    # Changing one resource price in TongHop_VT flows through to the total.
    old = xl.evaluate(f"TongHop_ChiPhi!F{row['TONG']}")
    wb = load_workbook(path)
    vt = wb["TongHop_VT"]
    xm = next(r for r in range(5, 60) if vt.cell(r, 1).value == "XM40")
    vt.cell(xm, 6).value += 1000
    wb.save(path)
    hao = next(r["haophi"] for r in e["resources"] if r["code"] == "XM40")
    expected = old + hao * 1000 * (1 + 0.065 + 0.011 + 0.025) * 1.055 * 1.08 * 1.05
    fresh = ExcelCompiler(filename=str(path))
    assert fresh.evaluate(f"TongHop_ChiPhi!F{row['TONG']}") == pytest.approx(expected, rel=1e-9)


def test_api_estimate(client):
    pid = client.post("/api/v1/boq/projects", json={"name": "Dự toán"}).json()["id"]
    client.post(f"/api/v1/boq/projects/{pid}/elements",
                json={"elements": [{"type": "COT", "n": 4, "hang_muc": "Nhà A"},
                                   {"type": "KHAC", "item_name": "Cửa", "unit": "bộ",
                                    "params": {"khoi_luong": 1}}]})
    assert client.get(f"/api/v1/boq/projects/{pid}/library").json()["norms"] == []
    assert client.post(f"/api/v1/boq/projects/{pid}/library/sample").json()["norms"] > 10
    e = client.get(f"/api/v1/boq/projects/{pid}/estimate").json()
    assert e["items"][0]["norm"] == "MAU.BT.COT" and e["items"][0]["wbs"] == "01.01.01"
    total = e["costs"][-1]["value"]
    r = client.put(f"/api/v1/boq/projects/{pid}/library/prices", json={"prices": {"XM40": 99999}})
    assert r.status_code == 200
    assert client.get(f"/api/v1/boq/projects/{pid}/estimate").json()["costs"][-1]["value"] > total
    assert client.put(f"/api/v1/boq/projects/{pid}/library/mapping",
                      json={"mapping": {"BT_COT|B22.5": "NOPE"}}).status_code == 400
    client.put(f"/api/v1/boq/projects/{pid}/library/mapping", json={"mapping": {"BT_COT": None}})
    e = client.get(f"/api/v1/boq/projects/{pid}/estimate").json()
    assert e["items"][0]["norm"] is None
    assert client.put(f"/api/v1/boq/projects/{pid}/rates", json={"rates": {"C": -1}}).status_code == 400
    client.put(f"/api/v1/boq/projects/{pid}/rates", json={"rates": {"VAT": 8}})
    assert client.get(f"/api/v1/boq/projects/{pid}/estimate").json()["rates"]["VAT"] == 8
    lib_x = client.get(f"/api/v1/boq/projects/{pid}/library.xlsx")
    r = client.post(f"/api/v1/boq/projects/{pid}/library/import", data={"mode": "replace"},
                    files={"file": ("dm.xlsx", lib_x.content, "application/octet-stream")})
    assert r.json()["errors"] == []
    x = client.get(f"/api/v1/boq/projects/{pid}/estimate.xlsx")
    assert x.status_code == 200 and x.content[:2] == b"PK"
    w = client.get(f"/api/v1/boq/projects/{pid}/wbs").json()
    assert w["nodes"][0] == {"code": "01", "level": 100, "name": "Nhà A", "unit": "", "qty": 0.0,
                             "cells": 3}
