"""BOQ engine: Bản vẽ → Shopdrawing → 3D → Thống kê → Bóc khối lượng → Excel.

Pure-Python, no database or web dependency, so every step can be tested alone.

Conventions
-----------
* All lengths are in metres (m). DXF import converts from drawing units.
* An element (cấu kiện) is one parametric object: type + dimensions + count.
  ``n`` copies are placed every ``dx``/``dy`` metres starting at (x, y, z).
* Every quantity line keeps an arithmetic expression (``expr``) built from the
  input numbers. The Excel export writes it as a live formula (``=expr``), so the
  checker can see exactly how each quantity was derived (diễn giải khối lượng).
"""
from __future__ import annotations

import ast
import io
import math
import operator
import re
from collections import OrderedDict
from dataclasses import dataclass, field

# ---------------------------------------------------------------------------
# Catalogue
# ---------------------------------------------------------------------------

# Work items (hạng mục công việc). Codes are internal; map them to norm codes
# (định mức) in the "Mã hiệu" column of the exported workbook if required.
WORK_ITEMS: "OrderedDict[str, dict]" = OrderedDict(
    [
        ("DAO", {"name": "Đào đất hố móng", "unit": "m3", "group": "Phần ngầm"}),
        ("BTL", {"name": "Bê tông lót móng", "unit": "m3", "group": "Phần ngầm"}),
        ("BT_MONG", {"name": "Bê tông móng", "unit": "m3", "group": "Phần ngầm"}),
        ("VK_MONG", {"name": "Ván khuôn móng", "unit": "m2", "group": "Phần ngầm"}),
        ("CT_MONG", {"name": "Cốt thép móng", "unit": "tấn", "group": "Phần ngầm"}),
        ("BT_COT", {"name": "Bê tông cột", "unit": "m3", "group": "Kết cấu"}),
        ("VK_COT", {"name": "Ván khuôn cột", "unit": "m2", "group": "Kết cấu"}),
        ("CT_COT", {"name": "Cốt thép cột", "unit": "tấn", "group": "Kết cấu"}),
        ("BT_DAM", {"name": "Bê tông dầm", "unit": "m3", "group": "Kết cấu"}),
        ("VK_DAM", {"name": "Ván khuôn dầm", "unit": "m2", "group": "Kết cấu"}),
        ("CT_DAM", {"name": "Cốt thép dầm", "unit": "tấn", "group": "Kết cấu"}),
        ("BT_SAN", {"name": "Bê tông sàn", "unit": "m3", "group": "Kết cấu"}),
        ("VK_SAN", {"name": "Ván khuôn sàn", "unit": "m2", "group": "Kết cấu"}),
        ("CT_SAN", {"name": "Cốt thép sàn", "unit": "tấn", "group": "Kết cấu"}),
        ("XAY", {"name": "Xây tường gạch", "unit": "m3", "group": "Kiến trúc"}),
        ("TRAT", {"name": "Trát tường", "unit": "m2", "group": "Hoàn thiện"}),
        ("SON", {"name": "Sơn tường", "unit": "m2", "group": "Hoàn thiện"}),
        ("LAT", {"name": "Lát nền", "unit": "m2", "group": "Hoàn thiện"}),
    ]
)

# Element types (loại cấu kiện) and their parameters: key -> (label, default).
ELEMENT_TYPES: "OrderedDict[str, dict]" = OrderedDict(
    [
        (
            "MONG",
            {
                "label": "Móng đơn",
                "color": "#8d6e63",
                "params": OrderedDict(
                    [
                        ("L", ("Dài L (m)", 1.5)),
                        ("B", ("Rộng B (m)", 1.5)),
                        ("H", ("Cao H (m)", 0.5)),
                        ("lot", ("Dày BT lót (m)", 0.1)),
                        ("sau_dao", ("Sâu đào (m)", 1.5)),
                        ("mo_rong", ("Mở rộng hố đào mỗi bên (m)", 0.3)),
                        ("ham_luong", ("Hàm lượng thép (kg/m3)", 80)),
                    ]
                ),
            },
        ),
        (
            "COT",
            {
                "label": "Cột",
                "color": "#546e7a",
                "params": OrderedDict(
                    [
                        ("b", ("Cạnh b (m)", 0.3)),
                        ("h", ("Cạnh h (m)", 0.3)),
                        ("H", ("Chiều cao H (m)", 3.6)),
                        ("ham_luong", ("Hàm lượng thép (kg/m3)", 150)),
                    ]
                ),
            },
        ),
        (
            "DAM",
            {
                "label": "Dầm",
                "color": "#1e88e5",
                "params": OrderedDict(
                    [
                        ("b", ("Rộng b (m)", 0.22)),
                        ("h", ("Cao h (m)", 0.5)),
                        ("L", ("Nhịp thông thủy L (m)", 6.0)),
                        ("t_san", ("Trừ dày sàn (m)", 0.12)),
                        ("goc", ("Hướng (0=X, 90=Y)", 0)),
                        ("ham_luong", ("Hàm lượng thép (kg/m3)", 130)),
                    ]
                ),
            },
        ),
        (
            "SAN",
            {
                "label": "Sàn",
                "color": "#90a4ae",
                "params": OrderedDict(
                    [
                        ("L", ("Dài L (m)", 6.0)),
                        ("B", ("Rộng B (m)", 6.0)),
                        ("t", ("Dày t (m)", 0.12)),
                        ("ham_luong", ("Hàm lượng thép (kg/m3)", 90)),
                    ]
                ),
            },
        ),
        (
            "TUONG",
            {
                "label": "Tường xây",
                "color": "#e57373",
                "params": OrderedDict(
                    [
                        ("L", ("Dài L (m)", 6.0)),
                        ("H", ("Cao H (m)", 3.2)),
                        ("t", ("Dày t (m)", 0.2)),
                        ("lo_cua", ("Diện tích lỗ cửa (m2)", 0)),
                        ("mat_trat", ("Số mặt trát/sơn", 2)),
                        ("goc", ("Hướng (0=X, 90=Y)", 0)),
                    ]
                ),
            },
        ),
        (
            "NEN",
            {
                "label": "Nền / Lát",
                "color": "#ffd54f",
                "params": OrderedDict(
                    [
                        ("L", ("Dài L (m)", 5.0)),
                        ("B", ("Rộng B (m)", 4.0)),
                    ]
                ),
            },
        ),
        (
            "KHAC",
            {
                "label": "Khác (nhập trực tiếp)",
                "color": "#ab47bc",
                "params": OrderedDict(
                    [
                        ("khoi_luong", ("Khối lượng / 1 cấu kiện", 0)),
                    ]
                ),
            },
        ),
    ]
)

DEFAULT_GRADE = "B22.5"


class BoqError(ValueError):
    """Raised for invalid user input (shown to the user as-is)."""


# ---------------------------------------------------------------------------
# Element model
# ---------------------------------------------------------------------------


@dataclass
class Element:
    type: str
    code: str = ""
    floor: str = ""
    n: int = 1
    x: float = 0.0
    y: float = 0.0
    z: float = 0.0
    dx: float = 0.0
    dy: float = 0.0
    params: dict = field(default_factory=dict)
    grade: str = ""  # mác bê tông / vật liệu
    # Only for type KHAC
    item_name: str = ""
    unit: str = ""
    id: int | None = None

    @classmethod
    def from_dict(cls, d: dict) -> "Element":
        etype = str(d.get("type", "")).strip().upper()
        if etype not in ELEMENT_TYPES:
            raise BoqError(f"Loại cấu kiện không hợp lệ: {d.get('type')!r}")
        spec = ELEMENT_TYPES[etype]["params"]
        raw = d.get("params") or {}
        params = {}
        for key, (label, default) in spec.items():
            value = raw.get(key, d.get(key, default))
            params[key] = _num(value, f"{label}")
            if key not in ("goc",) and params[key] < 0:
                raise BoqError(f"{label} không được âm")
        n_raw = d.get("n")
        n_val = _num(1 if n_raw in (None, "") else n_raw, "Số lượng")
        if n_val != int(n_val):
            raise BoqError("Số lượng phải là số nguyên")
        n = int(n_val)
        if n < 1:
            raise BoqError("Số lượng phải ≥ 1")
        return cls(
            type=etype,
            code=str(d.get("code") or "").strip() or etype,
            floor=str(d.get("floor") or "").strip(),
            n=n,
            x=_num(d.get("x", 0) or 0, "x"),
            y=_num(d.get("y", 0) or 0, "y"),
            z=_num(d.get("z", 0) or 0, "z"),
            dx=_num(d.get("dx", 0) or 0, "dx"),
            dy=_num(d.get("dy", 0) or 0, "dy"),
            params=params,
            grade=str(d.get("grade") or "").strip(),
            item_name=str(d.get("item_name") or "").strip(),
            unit=str(d.get("unit") or "").strip(),
            id=d.get("id"),
        )

    def to_dict(self) -> dict:
        return {
            "id": self.id,
            "type": self.type,
            "code": self.code,
            "floor": self.floor,
            "n": self.n,
            "x": self.x,
            "y": self.y,
            "z": self.z,
            "dx": self.dx,
            "dy": self.dy,
            "params": dict(self.params),
            "grade": self.grade,
            "item_name": self.item_name,
            "unit": self.unit,
        }


def _num(value, label: str) -> float:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        result = float(value)
    else:
        text = str(value).strip().replace(" ", "")
        # Accept Vietnamese decimal comma ("0,22") as well as "0.22".
        if text.count(",") == 1 and "." not in text:
            text = text.replace(",", ".")
        try:
            result = float(text)
        except ValueError:
            raise BoqError(f"{label}: '{value}' không phải là số") from None
    if not math.isfinite(result):
        raise BoqError(f"{label}: giá trị không hợp lệ")
    return result


# ---------------------------------------------------------------------------
# Arithmetic expressions (diễn giải)
# ---------------------------------------------------------------------------


def fmt(v: float) -> str:
    """Compact number for expressions: 0.30 -> '0.3', 6.0 -> '6'."""
    text = f"{v:.4f}".rstrip("0").rstrip(".")
    return text if text not in ("-0", "") else "0"


def mul(*parts) -> str:
    return "*".join(_wrap(p) for p in parts)


def _wrap(p) -> str:
    if isinstance(p, (int, float)):
        return fmt(p)
    text = str(p)
    return f"({text})" if re.search(r"[+\-]", text.lstrip("-")) else text


_OPS = {
    ast.Add: operator.add,
    ast.Sub: operator.sub,
    ast.Mult: operator.mul,
    ast.Div: operator.truediv,
    ast.USub: operator.neg,
    ast.UAdd: operator.pos,
}


def eval_expr(expr: str) -> float:
    """Safely evaluate a numeric expression (+ - * / parentheses only)."""

    def _eval(node):
        if isinstance(node, ast.Expression):
            return _eval(node.body)
        if isinstance(node, ast.Constant) and isinstance(node.value, (int, float)):
            return float(node.value)
        if isinstance(node, ast.BinOp) and type(node.op) in _OPS:
            return _OPS[type(node.op)](_eval(node.left), _eval(node.right))
        if isinstance(node, ast.UnaryOp) and type(node.op) in _OPS:
            return _OPS[type(node.op)](_eval(node.operand))
        raise BoqError(f"Biểu thức không hợp lệ: {expr}")

    return _eval(ast.parse(expr, mode="eval"))


# ---------------------------------------------------------------------------
# Step 5: Bóc khối lượng (quantity take-off)
# ---------------------------------------------------------------------------


@dataclass
class QtyLine:
    item: str  # work-item key (code|grade)
    code: str
    name: str
    unit: str
    element: str
    floor: str
    n: int
    expr: str  # full expression incl. the count, evaluates to qty
    note: str  # human-readable explanation
    qty: float = 0.0

    def to_dict(self) -> dict:
        return {
            "item": self.item,
            "code": self.code,
            "name": self.name,
            "unit": self.unit,
            "element": self.element,
            "floor": self.floor,
            "n": self.n,
            "expr": self.expr,
            "note": self.note,
            "qty": round(self.qty, 4),
        }


def _line(el: Element, code: str, per_unit_expr: str, note: str, scale: float = 1.0,
          grade: str | None = None) -> QtyLine:
    meta = WORK_ITEMS[code]
    name = meta["name"]
    key = code
    if grade:
        name = f"{name} {grade}"
        key = f"{code}|{grade}"
    expr = mul(el.n, per_unit_expr)
    if scale != 1.0:
        expr = f"{expr}/{fmt(1 / scale)}"
    line = QtyLine(
        item=key,
        code=code,
        name=name,
        unit=meta["unit"],
        element=el.code,
        floor=el.floor,
        n=el.n,
        expr=expr,
        note=note,
    )
    line.qty = eval_expr(expr)
    return line


def takeoff_element(el: Element) -> list[QtyLine]:
    p = el.params
    grade = el.grade or DEFAULT_GRADE
    lines: list[QtyLine] = []

    def rebar(code: str, concrete_expr: str):
        ratio = p.get("ham_luong", 0)
        if ratio > 0:
            lines.append(
                _line(el, code, mul(concrete_expr, ratio),
                      f"KL bê tông × {fmt(ratio)} kg/m3 (ước tính theo hàm lượng)",
                      scale=0.001)
            )

    if el.type == "MONG":
        L, B, H, lot = p["L"], p["B"], p["H"], p["lot"]
        a, depth = p["mo_rong"], p["sau_dao"]
        if depth > 0:
            lines.append(_line(
                el, "DAO", mul(f"{fmt(L)}+2*{fmt(a)}", f"{fmt(B)}+2*{fmt(a)}", depth),
                f"(L+2a)×(B+2a)×sâu = ({fmt(L)}+2×{fmt(a)})×({fmt(B)}+2×{fmt(a)})×{fmt(depth)}"))
        if lot > 0:
            lines.append(_line(
                el, "BTL", mul(f"{fmt(L)}+0.2", f"{fmt(B)}+0.2", lot),
                f"(L+0.2)×(B+0.2)×dày lót = ({fmt(L)}+0.2)×({fmt(B)}+0.2)×{fmt(lot)}"))
        concrete = mul(L, B, H)
        lines.append(_line(el, "BT_MONG", concrete,
                           f"L×B×H = {fmt(L)}×{fmt(B)}×{fmt(H)}", grade=grade))
        lines.append(_line(el, "VK_MONG", mul(f"2*({fmt(L)}+{fmt(B)})", H),
                           f"2×(L+B)×H = 2×({fmt(L)}+{fmt(B)})×{fmt(H)}"))
        rebar("CT_MONG", concrete)

    elif el.type == "COT":
        b, h, H = p["b"], p["h"], p["H"]
        concrete = mul(b, h, H)
        lines.append(_line(el, "BT_COT", concrete,
                           f"b×h×H = {fmt(b)}×{fmt(h)}×{fmt(H)}", grade=grade))
        lines.append(_line(el, "VK_COT", mul(f"2*({fmt(b)}+{fmt(h)})", H),
                           f"2×(b+h)×H = 2×({fmt(b)}+{fmt(h)})×{fmt(H)}"))
        rebar("CT_COT", concrete)

    elif el.type == "DAM":
        b, h, L, t = p["b"], p["h"], p["L"], p["t_san"]
        if t >= h:
            raise BoqError(f"{el.code}: dày sàn trừ ({fmt(t)}) phải nhỏ hơn chiều cao dầm ({fmt(h)})")
        hd = f"{fmt(h)}-{fmt(t)}" if t > 0 else fmt(h)
        concrete = mul(b, hd, L)
        lines.append(_line(el, "BT_DAM", concrete,
                           f"b×(h−t sàn)×L = {fmt(b)}×({hd})×{fmt(L)}", grade=grade))
        lines.append(_line(el, "VK_DAM", mul(f"{fmt(b)}+2*({hd})", L),
                           f"(đáy b + 2 thành (h−t))×L = ({fmt(b)}+2×({hd}))×{fmt(L)}"))
        rebar("CT_DAM", concrete)

    elif el.type == "SAN":
        L, B, t = p["L"], p["B"], p["t"]
        concrete = mul(L, B, t)
        lines.append(_line(el, "BT_SAN", concrete,
                           f"L×B×t = {fmt(L)}×{fmt(B)}×{fmt(t)}", grade=grade))
        lines.append(_line(el, "VK_SAN", mul(L, B), f"L×B = {fmt(L)}×{fmt(B)}"))
        rebar("CT_SAN", concrete)

    elif el.type == "TUONG":
        L, H, t, opening, faces = p["L"], p["H"], p["t"], p["lo_cua"], p["mat_trat"]
        if opening > L * H:
            raise BoqError(f"{el.code}: diện tích lỗ cửa lớn hơn diện tích tường")
        area = f"{fmt(L)}*{fmt(H)}-{fmt(opening)}" if opening > 0 else mul(L, H)
        area_note = f"(L×H−cửa) = ({fmt(L)}×{fmt(H)}−{fmt(opening)})" if opening > 0 \
            else f"L×H = {fmt(L)}×{fmt(H)}"
        lines.append(_line(el, "XAY", mul(area, t), f"{area_note}×t {fmt(t)}"))
        if faces > 0:
            lines.append(_line(el, "TRAT", mul(area, faces), f"{area_note}×{fmt(faces)} mặt"))
            lines.append(_line(el, "SON", mul(area, faces), f"{area_note}×{fmt(faces)} mặt"))

    elif el.type == "NEN":
        L, B = p["L"], p["B"]
        lines.append(_line(el, "LAT", mul(L, B), f"L×B = {fmt(L)}×{fmt(B)}"))

    elif el.type == "KHAC":
        q = p["khoi_luong"]
        name = el.item_name or el.code
        unit = el.unit or "cái"
        key = f"KHAC|{name}|{unit}"
        expr = mul(el.n, q)
        lines.append(QtyLine(item=key, code="KHAC", name=name, unit=unit,
                             element=el.code, floor=el.floor, n=el.n, expr=expr,
                             note="Nhập trực tiếp", qty=eval_expr(expr)))
    return lines


def takeoff(elements: list[Element]) -> list[QtyLine]:
    lines: list[QtyLine] = []
    for el in elements:
        lines.extend(takeoff_element(el))
    return lines


def summarize(lines: list[QtyLine], prices: dict | None = None) -> list[dict]:
    """Group take-off lines into BOQ items, ordered as in the catalogue."""
    prices = prices or {}
    order = {code: i for i, code in enumerate(WORK_ITEMS)}
    items: "OrderedDict[str, dict]" = OrderedDict()
    for ln in lines:
        it = items.setdefault(ln.item, {
            "item": ln.item, "code": ln.code, "name": ln.name, "unit": ln.unit,
            "group": WORK_ITEMS.get(ln.code, {}).get("group", "Khác"),
            "qty": 0.0, "lines": 0,
        })
        it["qty"] += ln.qty
        it["lines"] += 1
    result = sorted(items.values(), key=lambda it: (order.get(it["code"], 999), it["item"]))
    for it in result:
        it["qty"] = round(it["qty"], 4)
        price = float(prices.get(it["item"], 0) or 0)
        it["price"] = price
        it["amount"] = round(it["qty"] * price, 0)
    return result


# ---------------------------------------------------------------------------
# Step 4: Thống kê cấu kiện (element statistics)
# ---------------------------------------------------------------------------


def statistics(elements: list[Element], lines: list[QtyLine] | None = None) -> dict:
    lines = lines if lines is not None else takeoff(elements)
    by_type: "OrderedDict[str, dict]" = OrderedDict()
    for etype, meta in ELEMENT_TYPES.items():
        by_type[etype] = {"type": etype, "label": meta["label"], "groups": 0, "count": 0}
    floors: "OrderedDict[str, dict]" = OrderedDict()
    for el in elements:
        t = by_type[el.type]
        t["groups"] += 1
        t["count"] += el.n
        f = floors.setdefault(el.floor or "—", {"floor": el.floor or "—", "count": 0,
                                                "concrete": 0.0, "rebar": 0.0, "masonry": 0.0})
        f["count"] += el.n
    for ln in lines:
        f = floors.get(ln.floor or "—")
        if f is None:
            continue
        if ln.code.startswith("BT_"):
            f["concrete"] += ln.qty
        elif ln.code.startswith("CT_"):
            f["rebar"] += ln.qty
        elif ln.code == "XAY":
            f["masonry"] += ln.qty
    for f in floors.values():
        for k in ("concrete", "rebar", "masonry"):
            f[k] = round(f[k], 3)
    totals = {
        "elements": sum(el.n for el in elements),
        "groups": len(elements),
        "concrete": round(sum(ln.qty for ln in lines if ln.code.startswith("BT_")), 3),
        "rebar": round(sum(ln.qty for ln in lines if ln.code.startswith("CT_")), 3),
        "formwork": round(sum(ln.qty for ln in lines if ln.code.startswith("VK_")), 3),
        "masonry": round(sum(ln.qty for ln in lines if ln.code == "XAY"), 3),
    }
    return {
        "by_type": [t for t in by_type.values() if t["count"]],
        "by_floor": list(floors.values()),
        "totals": totals,
    }


# ---------------------------------------------------------------------------
# Step 3: 3D model (axis-aligned / rotated boxes for a WebGL viewer)
# ---------------------------------------------------------------------------


def instances(el: Element):
    for i in range(el.n):
        yield el.x + i * el.dx, el.y + i * el.dy


def model3d(elements: list[Element]) -> dict:
    """Boxes in metres: centre (cx, cy, cz), size (sx, sy, sz), rotation about Z.

    z is up. x/y are plan coordinates. Origin meaning per type:
    * COT, MONG: (x, y) = centre of the element.
    * DAM, TUONG: (x, y) = start of the axis, running along ``goc`` degrees.
    * SAN, NEN: (x, y) = lower-left corner.
    """
    boxes = []
    for el in elements:
        p = el.params
        color = ELEMENT_TYPES[el.type]["color"]
        for k, (x, y) in enumerate(instances(el)):
            label = el.code if el.n == 1 else f"{el.code} #{k + 1}"
            box = None
            if el.type == "COT":
                box = (x, y, el.z + p["H"] / 2, p["b"], p["h"], p["H"], 0)
            elif el.type == "MONG":
                box = (x, y, el.z + p["H"] / 2, p["L"], p["B"], p["H"], 0)
            elif el.type in ("DAM", "TUONG"):
                ang = math.radians(p.get("goc", 0))
                L = p["L"]
                width = p["b"] if el.type == "DAM" else p["t"]
                height = p["h"] if el.type == "DAM" else p["H"]
                cx = x + math.cos(ang) * L / 2
                cy = y + math.sin(ang) * L / 2
                box = (cx, cy, el.z + height / 2, L, width, height, p.get("goc", 0))
            elif el.type == "SAN":
                box = (x + p["L"] / 2, y + p["B"] / 2, el.z + p["t"] / 2, p["L"], p["B"], p["t"], 0)
            elif el.type == "NEN":
                box = (x + p["L"] / 2, y + p["B"] / 2, el.z + 0.01, p["L"], p["B"], 0.02, 0)
            if box is None:
                continue
            cx, cy, cz, sx, sy, sz, rot = box
            boxes.append({
                "id": el.id, "label": label, "type": el.type, "floor": el.floor,
                "color": color,
                "c": [round(cx, 4), round(cy, 4), round(cz, 4)],
                "s": [round(sx, 4), round(sy, 4), round(sz, 4)],
                "rot": rot,
            })
    if boxes:
        xs = [b["c"][0] for b in boxes]
        ys = [b["c"][1] for b in boxes]
        zs = [b["c"][2] for b in boxes]
        bounds = {"min": [min(xs), min(ys), min(zs)], "max": [max(xs), max(ys), max(zs)]}
    else:
        bounds = {"min": [0, 0, 0], "max": [0, 0, 0]}
    return {"boxes": boxes, "bounds": bounds}


# ---------------------------------------------------------------------------
# Step 2: Shopdrawing (detail sketch per element, SVG with dimensions)
# ---------------------------------------------------------------------------


def _esc(text) -> str:
    return (str(text).replace("&", "&amp;").replace("<", "&lt;")
            .replace(">", "&gt;").replace('"', "&quot;"))


def _dim_h(x1, x2, y, text):
    return (f'<line x1="{x1:.1f}" y1="{y:.1f}" x2="{x2:.1f}" y2="{y:.1f}" class="dim"/>'
            f'<line x1="{x1:.1f}" y1="{y - 5:.1f}" x2="{x1:.1f}" y2="{y + 5:.1f}" class="dim"/>'
            f'<line x1="{x2:.1f}" y1="{y - 5:.1f}" x2="{x2:.1f}" y2="{y + 5:.1f}" class="dim"/>'
            f'<text x="{(x1 + x2) / 2:.1f}" y="{y - 6:.1f}" text-anchor="middle">{_esc(text)}</text>')


def _dim_v(x, y1, y2, text):
    cy = (y1 + y2) / 2
    return (f'<line x1="{x:.1f}" y1="{y1:.1f}" x2="{x:.1f}" y2="{y2:.1f}" class="dim"/>'
            f'<line x1="{x - 5:.1f}" y1="{y1:.1f}" x2="{x + 5:.1f}" y2="{y1:.1f}" class="dim"/>'
            f'<line x1="{x - 5:.1f}" y1="{y2:.1f}" x2="{x + 5:.1f}" y2="{y2:.1f}" class="dim"/>'
            f'<text x="{x - 8:.1f}" y="{cy:.1f}" text-anchor="middle" '
            f'transform="rotate(-90 {x - 8:.1f} {cy:.1f})">{_esc(text)}</text>')


def _mm(v: float) -> str:
    return f"{round(v * 1000):d}"


def shop_svg(el: Element) -> str:
    """Two views per element (plan/section + elevation) with dimensions in mm."""
    p = el.params
    W, Hc = 640, 300
    views = []  # (title, width_m, height_m, labels(w,h), extra)
    if el.type == "COT":
        views = [("MẶT CẮT", p["b"], p["h"], ("b", "h")), ("MẶT ĐỨNG", p["b"], p["H"], ("b", "H"))]
    elif el.type == "MONG":
        views = [("MẶT BẰNG", p["L"], p["B"], ("L", "B")), ("MẶT CẮT", p["L"], p["H"], ("L", "H"))]
    elif el.type == "DAM":
        views = [("MẶT CẮT", p["b"], p["h"], ("b", "h")), ("MẶT ĐỨNG", p["L"], p["h"], ("L", "h"))]
    elif el.type == "SAN":
        views = [("MẶT BẰNG", p["L"], p["B"], ("L", "B")), ("MẶT CẮT", p["L"], p["t"], ("L", "t"))]
    elif el.type == "TUONG":
        views = [("MẶT ĐỨNG", p["L"], p["H"], ("L", "H")), ("MẶT CẮT", p["t"], p["H"], ("t", "H"))]
    elif el.type == "NEN":
        views = [("MẶT BẰNG", p["L"], p["B"], ("L", "B"))]

    parts = [
        f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {W} {Hc + 60}" '
        f'font-family="Arial, sans-serif" font-size="12">',
        "<style>.o{fill:#eef3f8;stroke:#1f2937;stroke-width:2}.dim{stroke:#c62828;stroke-width:1}"
        "text{fill:#1f2937}.t{font-weight:bold;font-size:13px}.h{font-weight:bold;font-size:15px}"
        ".hatch{fill:url(#hatch)}</style>",
        '<defs><pattern id="hatch" width="8" height="8" patternUnits="userSpaceOnUse" '
        'patternTransform="rotate(45)"><line x1="0" y1="0" x2="0" y2="8" stroke="#90a4ae" '
        'stroke-width="1"/></pattern></defs>',
        f'<rect x="0" y="0" width="{W}" height="{Hc + 60}" fill="#ffffff"/>',
        f'<text x="12" y="22" class="h">SHOPDRAWING · {_esc(el.code)} · '
        f'{_esc(ELEMENT_TYPES[el.type]["label"])} · SL {el.n}'
        f'{" · Tầng " + _esc(el.floor) if el.floor else ""}</text>',
    ]
    if not views:
        parts.append('<text x="12" y="60">Cấu kiện nhập khối lượng trực tiếp — không có bản vẽ chi tiết.</text>')
    else:
        slot = W / len(views)
        for i, (title, wm, hm, (lw, lh)) in enumerate(views):
            avail_w, avail_h = slot - 110, Hc - 130
            scale = min(avail_w / max(wm, 1e-6), avail_h / max(hm, 1e-6))
            dw, dh = max(wm * scale, 2), max(hm * scale, 2)
            ox = i * slot + 70 + (avail_w - dw) / 2
            oy = 70 + (avail_h - dh) / 2
            cls = "o hatch" if title == "MẶT CẮT" else "o"
            parts.append(f'<text x="{i * slot + slot / 2:.1f}" y="{Hc - 22}" text-anchor="middle" '
                         f'class="t">{title}</text>')
            parts.append(f'<rect x="{ox:.1f}" y="{oy:.1f}" width="{dw:.1f}" height="{dh:.1f}" class="{cls}"/>')
            parts.append(_dim_h(ox, ox + dw, oy - 14, f"{lw} = {_mm(wm)}"))
            parts.append(_dim_v(ox - 16, oy, oy + dh, f"{lh} = {_mm(hm)}"))
    lines = takeoff_element(el)
    y = Hc + 6
    parts.append(f'<line x1="0" y1="{y - 14}" x2="{W}" y2="{y - 14}" stroke="#cfd8dc"/>')
    summary = " · ".join(f"{ln.name}: {ln.qty:.3f} {ln.unit}" for ln in lines[:4])
    parts.append(f'<text x="12" y="{y + 4}">{_esc(summary)}</text>')
    grade = el.grade or (DEFAULT_GRADE if el.type in ("MONG", "COT", "DAM", "SAN") else "")
    info = ["Kích thước: mm"]
    if grade:
        info.append(f"Bê tông {_esc(grade)}")
    if "ham_luong" in p:
        info.append(f"Hàm lượng thép: {fmt(p['ham_luong'])} kg/m3")
    parts.append(f'<text x="12" y="{y + 24}">{" · ".join(info)}</text>')
    parts.append("</svg>")
    return "".join(parts)


# ---------------------------------------------------------------------------
# Step 1: Bản vẽ — import from DXF (CAD) or an Excel template
# ---------------------------------------------------------------------------

# Layer-name keywords → element type. Matched case-insensitively, first match wins.
LAYER_RULES = [
    (("MONG", "FOOT", "FOUND"), "MONG"),
    (("COT", "COL"), "COT"),
    (("DAM", "BEAM"), "DAM"),
    (("SAN", "SLAB"), "SAN"),
    (("TUONG", "WALL"), "TUONG"),
    (("NEN", "FLOOR", "LAT"), "NEN"),
]


def _strip_accents(text: str) -> str:
    import unicodedata
    text = unicodedata.normalize("NFD", text)
    text = "".join(c for c in text if unicodedata.category(c) != "Mn")
    return text.replace("đ", "d").replace("Đ", "D")


def layer_type(layer: str) -> str | None:
    name = _strip_accents(layer).upper()
    for keys, etype in LAYER_RULES:
        if any(k in name for k in keys):
            return etype
    return None


def import_dxf(data: bytes, unit_scale: float = 0.001, defaults: dict | None = None,
               floor: str = "", z: float = 0.0) -> dict:
    """Read closed polylines from a DXF and turn them into element candidates.

    * Layer name decides the element type (see LAYER_RULES), e.g. "KC-COT", "BEAM".
    * Each closed outline is reduced to its bounding rectangle in plan.
      Long side → length, short side → width. Heights come from ``defaults``.
    * ``unit_scale`` converts drawing units to metres (mm drawing → 0.001).
    Identical outlines on the same layer are merged into one element with ``n``.
    """
    import ezdxf
    from ezdxf import recover

    defaults = defaults or {}
    try:
        doc, _ = recover.read(io.BytesIO(data))
    except Exception:
        try:
            doc = ezdxf.read(io.StringIO(data.decode("utf-8", errors="ignore")))
        except Exception as exc:
            raise BoqError(f"Không đọc được file DXF: {exc}") from None

    skipped: dict[str, int] = {}
    merged: "OrderedDict[tuple, dict]" = OrderedDict()
    counters: dict[str, int] = {}

    for e in doc.modelspace():
        kind = e.dxftype()
        if kind == "LWPOLYLINE":
            if not (e.closed or _closes(list(e.get_points("xy")))):
                continue
            pts = [(float(px), float(py)) for px, py in e.get_points("xy")]
        elif kind == "POLYLINE" and not e.is_3d_polyline:
            pts = [(float(v.dxf.location.x), float(v.dxf.location.y)) for v in e.vertices]
            if not (e.is_closed or _closes(pts)):
                continue
        else:
            continue
        layer = e.dxf.layer
        etype = layer_type(layer)
        if etype is None or len(pts) < 3:
            skipped[layer] = skipped.get(layer, 0) + 1
            continue
        xs = [px * unit_scale for px, _ in pts]
        ys = [py * unit_scale for _, py in pts]
        w, d = max(xs) - min(xs), max(ys) - min(ys)
        if w <= 0 or d <= 0:
            continue
        x0, y0 = min(xs), min(ys)
        long_, short = max(w, d), min(w, d)
        storey = defaults.get("H_cot", 3.6)
        ez = z  # base elevation of this element: beams/slabs sit flush with the column top
        along_y = d > w
        params: dict = {}
        x, y = x0, y0
        if etype == "COT":
            params = {"b": round(w, 3), "h": round(d, 3), "H": defaults.get("H_cot", 3.6)}
            x, y = x0 + w / 2, y0 + d / 2
        elif etype == "MONG":
            params = {"L": round(w, 3), "B": round(d, 3), "H": defaults.get("H_mong", 0.5)}
            x, y = x0 + w / 2, y0 + d / 2
            ez = z - params["H"]
        elif etype == "DAM":
            params = {"L": round(long_, 3), "b": round(short, 3), "h": defaults.get("h_dam", 0.5),
                      "t_san": defaults.get("t_san", 0.12), "goc": 90 if along_y else 0}
            x, y = (x0 + w / 2, y0) if along_y else (x0, y0 + d / 2)
            ez = z + storey - params["h"]
        elif etype == "TUONG":
            params = {"L": round(long_, 3), "t": round(short, 3), "H": defaults.get("H_tuong", 3.2),
                      "goc": 90 if along_y else 0}
            x, y = (x0 + w / 2, y0) if along_y else (x0, y0 + d / 2)
        elif etype == "SAN":
            area = abs(_shoelace(list(zip(xs, ys))))
            if abs(area - w * d) > 0.01 * w * d:
                # Non-rectangular slab: keep the true area, express as equivalent L×B.
                params = {"L": round(w, 3), "B": round(area / w, 3), "t": defaults.get("t_san", 0.12)}
            else:
                params = {"L": round(w, 3), "B": round(d, 3), "t": defaults.get("t_san", 0.12)}
            ez = z + storey - params["t"]
        elif etype == "NEN":
            area = abs(_shoelace(list(zip(xs, ys))))
            params = {"L": round(w, 3), "B": round(area / w, 3)}
        sig = (etype, layer, ez) + tuple(sorted((k, round(v, 3)) for k, v in params.items()))
        if sig in merged:
            merged[sig]["n"] += 1
            merged[sig]["_positions"].append((round(x, 3), round(y, 3)))
            continue
        counters[etype] = counters.get(etype, 0) + 1
        prefix = {"COT": "C", "MONG": "M", "DAM": "D", "SAN": "S", "TUONG": "T", "NEN": "N"}[etype]
        merged[sig] = {
            "type": etype, "code": f"{prefix}{counters[etype]}", "floor": floor, "n": 1,
            "x": round(x, 3), "y": round(y, 3), "z": round(ez, 3), "params": params,
            "source": f"DXF layer {layer}", "_positions": [(round(x, 3), round(y, 3))],
        }

    elements = []
    for cand in merged.values():
        positions = cand.pop("_positions")
        # Evenly spaced copies along a line → keep spacing so the 3D model is right.
        if len(positions) > 1:
            positions.sort()
            sx = positions[1][0] - positions[0][0]
            sy = positions[1][1] - positions[0][1]
            regular = all(
                abs(positions[i][0] - positions[0][0] - i * sx) < 1e-3
                and abs(positions[i][1] - positions[0][1] - i * sy) < 1e-3
                for i in range(len(positions))
            )
            if regular:
                cand["x"], cand["y"] = positions[0]
                cand["dx"], cand["dy"] = round(sx, 3), round(sy, 3)
            else:
                cand["positions"] = positions
        elements.append(cand)

    # Irregular groups: split into one element per position so 3D stays faithful.
    out = []
    for cand in elements:
        positions = cand.pop("positions", None)
        if positions:
            for px, py in positions:
                out.append({**cand, "n": 1, "x": px, "y": py})
        else:
            out.append(cand)
    # Normalise through Element so every parameter (incl. defaults) is present.
    normalised = []
    for cand in out:
        full = Element.from_dict(cand).to_dict()
        full["source"] = cand["source"]
        normalised.append(full)
    return {"elements": normalised, "skipped_layers": skipped}


def _closes(pts) -> bool:
    return len(pts) > 2 and math.dist(pts[0], pts[-1]) < 1e-6


def _shoelace(pts) -> float:
    s = 0.0
    for (x1, y1), (x2, y2) in zip(pts, pts[1:] + pts[:1]):
        s += x1 * y2 - x2 * y1
    return s / 2


TEMPLATE_COLUMNS = ["type", "code", "floor", "n", "x", "y", "z", "dx", "dy", "grade",
                    "p1", "p2", "p3", "p4", "p5", "p6", "item_name", "unit"]


def template_workbook() -> bytes:
    """Excel template users fill in to bulk-enter elements (step 1 shortcut)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "CauKien"
    headers = ["Loại", "Mã CK", "Tầng", "Số lượng", "X (m)", "Y (m)", "Z (m)",
               "Bước X (m)", "Bước Y (m)", "Mác BT", "TS1", "TS2", "TS3", "TS4", "TS5", "TS6",
               "Tên công việc (Khác)", "Đơn vị (Khác)"]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F2937")
    examples = [
        ["MONG", "M1", "Móng", 6, 0, 0, -1.5, 6, 0, "B25", 1.8, 1.8, 0.6, 0.1, 1.6, 0.3, None, None],
        ["COT", "C1", "L01", 6, 0, 0, 0, 6, 0, "B25", 0.3, 0.3, 3.6, 150, None, None, None, None],
        ["DAM", "D1", "L01", 5, 0, 0, 3.1, 6, 0, "B25", 0.22, 0.5, 6, 0.12, 0, 130, None, None],
        ["SAN", "S1", "L01", 1, 0, -3, 3.48, 0, 0, "B25", 30, 6, 0.12, 90, None, None, None, None],
        ["TUONG", "T1", "L01", 2, 0, -3, 0, 0, 6, None, 30, 3.1, 0.2, 6, 2, 0, None, None],
        ["NEN", "N1", "L01", 1, 0, -3, 0, 0, 0, None, 30, 6, None, None, None, None, None, None],
        ["KHAC", "K1", "L01", 4, 0, 0, 0, 0, 0, None, 1, None, None, None, None, None,
         "Cửa đi gỗ D1 900×2200", "bộ"],
    ]
    for row in examples:
        ws.append(row)
    dv = DataValidation(type="list", formula1='"' + ",".join(ELEMENT_TYPES) + '"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("A2:A1000")
    for col, width in zip("ABCDEFGHIJKLMNOPQR", [9, 8, 8, 9, 7, 7, 7, 9, 9, 8, 7, 7, 7, 7, 7, 7, 26, 12]):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"

    guide = wb.create_sheet("HuongDan")
    guide.append(["Loại", "Tên", "TS1", "TS2", "TS3", "TS4", "TS5", "TS6"])
    for c in guide[1]:
        c.font = Font(bold=True)
    for etype, meta in ELEMENT_TYPES.items():
        guide.append([etype, meta["label"]] + [f"{label} [mặc định {default}]"
                                               for label, default in meta["params"].values()])
    guide.append([])
    guide.append(["Ghi chú: đơn vị mét. Ô TS để trống sẽ dùng giá trị mặc định. "
                  "Số lượng > 1 với Bước X/Y sẽ rải cấu kiện đều trên mô hình 3D."])
    guide.column_dimensions["B"].width = 22
    for col in "CDEFGH":
        guide.column_dimensions[col].width = 30
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def import_template(data: bytes) -> dict:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:
        raise BoqError(f"Không đọc được file Excel: {exc}") from None
    ws = wb["CauKien"] if "CauKien" in wb.sheetnames else wb.active
    elements, errors = [], []
    for r, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row or row[0] in (None, ""):
            continue
        row = list(row) + [None] * (len(TEMPLATE_COLUMNS) - len(row))
        rec = dict(zip(TEMPLATE_COLUMNS, row))
        etype = str(rec["type"]).strip().upper()
        if etype not in ELEMENT_TYPES:
            errors.append(f"Dòng {r}: loại '{rec['type']}' không hợp lệ")
            continue
        keys = list(ELEMENT_TYPES[etype]["params"])
        params = {}
        for i, key in enumerate(keys):
            v = rec.get(f"p{i + 1}")
            if v not in (None, ""):
                params[key] = v
        d = {k: rec[k] for k in ("code", "floor", "n", "x", "y", "z", "dx", "dy", "grade",
                                 "item_name", "unit") if rec.get(k) not in (None, "")}
        d["type"] = etype
        d["params"] = params
        try:
            elements.append(Element.from_dict(d).to_dict())
        except BoqError as exc:
            errors.append(f"Dòng {r}: {exc}")
    return {"elements": elements, "errors": errors}


# ---------------------------------------------------------------------------
# Step 6: Excel export
# ---------------------------------------------------------------------------


def export_workbook(project: dict, elements: list[Element], prices: dict | None = None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    prices = prices or {}
    lines = takeoff(elements)
    summary = summarize(lines, prices)
    stats = statistics(elements, lines)

    thin = Side(style="thin", color="9CA3AF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="1F2937")
    group_fill = PatternFill("solid", fgColor="E5E7EB")
    head_font = Font(bold=True, color="FFFFFF")
    wrap = Alignment(wrap_text=True, vertical="center")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    qty_fmt = "#,##0.000"
    money_fmt = "#,##0"

    def header(ws, row, titles, widths=None):
        for i, t in enumerate(titles, start=1):
            c = ws.cell(row=row, column=i, value=t)
            c.font, c.fill, c.alignment, c.border = head_font, head_fill, center, border
            if widths:
                ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
        ws.row_dimensions[row].height = 30

    def title(ws, text, ncols):
        ws.cell(row=1, column=1, value=text).font = Font(bold=True, size=14)
        ws.cell(row=2, column=1,
                value=f"Công trình: {project.get('name', '')}   ·   Mã: {project.get('code', '')}"
                      f"   ·   Địa điểm: {project.get('location', '') or ''}").font = Font(italic=True)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)

    wb = Workbook()

    # --- Sheet 1: BOQ summary (KL linked to detail sheet by SUMIF) ---------
    ws = wb.active
    ws.title = "TongHop_BOQ"
    title(ws, "BẢNG TỔNG HỢP KHỐI LƯỢNG (BOQ)", 8)
    header(ws, 4, ["STT", "Mã hiệu", "Nội dung công việc", "Đơn vị", "Khối lượng",
                   "Đơn giá (VNĐ)", "Thành tiền (VNĐ)", "Mã nội bộ"],
           [6, 14, 42, 9, 15, 16, 18, 22])
    row = 5
    first_item_row = row
    group = None
    stt = 0
    for it in summary:
        if it["group"] != group:
            group = it["group"]
            c = ws.cell(row=row, column=3, value=group.upper())
            c.font = Font(bold=True)
            for col in range(1, 9):
                ws.cell(row=row, column=col).fill = group_fill
                ws.cell(row=row, column=col).border = border
            row += 1
        stt += 1
        ws.cell(row=row, column=1, value=stt).alignment = center
        ws.cell(row=row, column=2, value="")  # Mã hiệu định mức: người dùng điền
        ws.cell(row=row, column=3, value=it["name"]).alignment = wrap
        ws.cell(row=row, column=4, value=it["unit"]).alignment = center
        ws.cell(row=row, column=5,
                value=f'=SUMIF(ChiTiet!$K:$K,$H{row},ChiTiet!$J:$J)').number_format = qty_fmt
        ws.cell(row=row, column=6, value=it["price"] or None).number_format = money_fmt
        ws.cell(row=row, column=7, value=f"=E{row}*F{row}").number_format = money_fmt
        ws.cell(row=row, column=8, value=it["item"]).font = Font(color="6B7280", size=9)
        for col in range(1, 9):
            ws.cell(row=row, column=col).border = border
        row += 1
    last_item_row = row - 1
    ws.cell(row=row, column=3, value="TỔNG CỘNG").font = Font(bold=True)
    total = ws.cell(row=row, column=7,
                    value=f"=SUM(G{first_item_row}:G{last_item_row})" if summary else 0)
    total.font, total.number_format = Font(bold=True), money_fmt
    for col in range(1, 9):
        ws.cell(row=row, column=col).border = border
    ws.cell(row=row + 2, column=1,
            value="Ghi chú: Khối lượng liên kết công thức với sheet ChiTiet. "
                  "Cốt thép ước tính theo hàm lượng kg/m3 — thay bằng bảng thống kê thép "
                  "shopdrawing khi có.").font = Font(italic=True, color="6B7280")
    ws.freeze_panes = "A5"

    # --- Sheet 2: detailed take-off with live formulas ---------------------
    wd = wb.create_sheet("ChiTiet")
    title(wd, "BẢNG TÍNH CHI TIẾT KHỐI LƯỢNG (DIỄN GIẢI)", 11)
    header(wd, 4, ["STT", "Nội dung công việc", "Cấu kiện", "Tầng", "SL", "Diễn giải",
                   "Công thức", "Đơn vị", "KL 1 CK", "Khối lượng", "Mã nội bộ"],
           [6, 30, 10, 8, 6, 46, 30, 8, 12, 14, 22])
    r = 5
    for i, ln in enumerate(lines, start=1):
        wd.cell(row=r, column=1, value=i).alignment = center
        wd.cell(row=r, column=2, value=ln.name).alignment = wrap
        wd.cell(row=r, column=3, value=ln.element)
        wd.cell(row=r, column=4, value=ln.floor)
        wd.cell(row=r, column=5, value=ln.n).alignment = center
        wd.cell(row=r, column=6, value=ln.note).alignment = wrap
        wd.cell(row=r, column=7, value=ln.expr).font = Font(color="6B7280", size=9)
        wd.cell(row=r, column=8, value=ln.unit).alignment = center
        wd.cell(row=r, column=9, value=f"=J{r}/E{r}").number_format = qty_fmt
        wd.cell(row=r, column=10, value=f"={ln.expr}").number_format = qty_fmt
        wd.cell(row=r, column=11, value=ln.item).font = Font(color="6B7280", size=9)
        for col in range(1, 12):
            wd.cell(row=r, column=col).border = border
        r += 1
    wd.freeze_panes = "A5"
    wd.auto_filter.ref = f"A4:K{max(r - 1, 4)}"

    # --- Sheet 3: element statistics --------------------------------------
    wt = wb.create_sheet("ThongKe_CauKien")
    title(wt, "BẢNG THỐNG KÊ CẤU KIỆN", 9)
    header(wt, 4, ["STT", "Mã CK", "Loại", "Tầng", "Số lượng", "Kích thước (m)", "Mác",
                   "BT 1 CK (m3)", "BT tổng (m3)"], [6, 10, 16, 8, 10, 46, 9, 13, 13])
    r = 5
    for i, el in enumerate(elements, start=1):
        spec = ELEMENT_TYPES[el.type]["params"]
        dims = ", ".join(f"{k}={fmt(el.params[k])}" for k in spec if k not in ("ham_luong", "goc"))
        conc = sum(ln.qty for ln in takeoff_element(el) if ln.code.startswith("BT_"))
        vals = [i, el.code, ELEMENT_TYPES[el.type]["label"], el.floor, el.n, dims,
                el.grade or (DEFAULT_GRADE if conc else ""), round(conc / el.n, 4) if conc else None,
                f"=E{r}*H{r}" if conc else None]
        for col, v in enumerate(vals, start=1):
            c = wt.cell(row=r, column=col, value=v)
            c.border = border
            if col in (8, 9):
                c.number_format = qty_fmt
        r += 1
    r += 1
    wt.cell(row=r, column=2, value="TỔNG HỢP THEO TẦNG").font = Font(bold=True)
    r += 1
    header(wt, r, ["", "Tầng", "Số CK", "BT (m3)", "Thép (tấn)", "Xây (m3)"])
    r += 1
    for f in stats["by_floor"]:
        for col, v in enumerate(["", f["floor"], f["count"], f["concrete"], f["rebar"], f["masonry"]],
                                start=1):
            c = wt.cell(row=r, column=col, value=v)
            c.border = border
        r += 1
    wt.freeze_panes = "A5"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
